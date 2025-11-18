import mimetypes
import os
from typing import Literal, Union
from dotenv import load_dotenv
import boto3
import io
from pydantic import BaseModel
import requests
from database import Functions
from urllib import parse as url_parser
from PyPDF2.generic import DictionaryObject
from PyPDF2 import PdfReader
from docx import Document
import openpyxl
from pptx import Presentation

from helpers import Helpers

load_dotenv()
db_func = Functions()

class DummyFile(BaseModel):
    content: bytes
    size: int
    name: str
    ext: str
    content_type: str 


class BucketConsole:
    all_formnames = db_func.get_all_formnames()
    AWS_ACCESS_ID = os.getenv("AWS_ACCESS_ID")
    AWS_SECRET_KEY = os.getenv("AWS_SECRET_KEY")
    bucket_name = "doloreschatbucket"
    session = boto3.Session(
        aws_access_key_id=AWS_ACCESS_ID,
        aws_secret_access_key=AWS_SECRET_KEY,
    )
    region = session.region_name
    s3_client = session.client('s3')
    s3_resource = session.resource("s3")
    bucket = s3_resource.Bucket(bucket_name)
    base_url = f"https://{bucket_name}.s3.{region}.amazonaws.com"
    image_extensions = ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', 'webp', '.svg']
    allowed_case_extensions = {'.pdf', '.docx', '.doc', '.txt', '.jpg', '.jpeg', '.png', '.heic'}

    def upload_file(self, file: DummyFile, bucket_directory: str, is_previewable: bool = True):
        bucket_directory = bucket_directory.replace(" ", "+")
        file_path = f"./tmp/{file.name}.{file.ext}"
        with open(file_path, "wb") as fp:
            fp.write(file.content)
        if not is_previewable:
            self.s3_client.upload_file(
                Filename=file_path,
                Bucket=self.bucket_name,
                Key=bucket_directory
            )
            return f"{self.base_url}/{bucket_directory}"
        with open(file_path, "rb") as stream:
            self.s3_client.upload_fileobj(
                FileObj=stream,
                Bucket=self.bucket_name,
                Key=bucket_directory,
                ExtraArgs={"ContentType": file.content_type, "ACL": "public-read"},
            )

    def upload_new_profile_picture(self, file_data: bytes, user_type: Literal["immigrants", "lawyers", "nonlawyers", "lawstudents", "paralegals"], username: str) -> str:
        file = DummyFile(
            content = file_data,
            size = len(file_data),
            name = "user",
            ext = "png",
            content_type = "image/png"
        )
        bucket_directory = f"{user_type}/{username}/profile_pic/user.png"
        self.upload_file(file, bucket_directory, is_previewable=True)

    def _extract_text_content(self, filename: str, file_content: bytes) -> str:
        file_ext = filename.split(".")[-1].lower()
        try:
            if file_ext in ["txt", "csv", "json"]:
                return file_content.decode("utf-8")
            elif file_ext == "pdf":
                return self._extract_pdf_text(file_content)
            elif file_ext == "docx":
                return self._extract_docx_text(file_content)
            elif file_ext == "xlsx":
                return self._extract_xlsx_text(file_content)
            elif file_ext == "pptx":
                return self._extract_pptx_text(file_content)
            else:
                return None
        except:
            return ""

    def _extract_pdf_text(self, file_content: bytes) -> str:
        text = ""
        try:
            pdf = PdfReader(io.BytesIO(file_content))
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        except Exception as e:
            print(f"Error extracting text from PDF page: {e}")
        return text

    def _extract_docx_text(self, file_content: bytes) -> str:
        doc = Document(io.BytesIO(file_content))
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])

    def _extract_xlsx_text(self, file_content: bytes) -> str:
        workbook = openpyxl.load_workbook(io.BytesIO(file_content))
        text = []
        for sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
            for row in worksheet.iter_rows(values_only=True):
                text.append("\t".join(str(cell) for cell in row if cell is not None))
        return "\n".join(text)

    def _extract_pptx_text(self, file_content: bytes) -> str:
        prs = Presentation(io.BytesIO(file_content))
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)

    def get_object_head(self, s3_key: str) -> bool:
        try:
            self.s3_client.head_object(Bucket=self.bucket_name, Key=s3_key)
            return True
        except:
            return False

    def crosscheck_existing(
        self, destination_path: str, filename: str
    ) -> tuple[str, str]:
        counter = 1
        s3_key = f"{destination_path}/{filename}"
        stem, ext = os.path.splitext(filename)
        while self.get_object_head(s3_key):
            filename = f"{stem}_{counter}{ext}"
            counter += 1
            s3_key = f"{destination_path}/{filename}"
        return s3_key, filename

    def download_file(self, file_url: str):
        parsed = url_parser.urlparse(file_url)
        path = parsed.path
        file_key = path.lstrip("/")
        filename = file_key.split("/")[-1]
        file_contents = {}

        if filename and "blob" not in filename:
            try:
                file_obj = self.s3_client.get_object(Bucket=self.bucket_name, Key=file_key)
                file_content = file_obj["Body"].read()
                file_contents["filename"] = filename
                if filename.lower().endswith(".pdf"):
                    if self._is_pdf_image(file_content):
                        file_contents["image_data"] = file_content
                        file_contents["content_type"] = "pdf_images"
                    else:
                        text_content = self._extract_pdf_text(file_content)
                        if text_content is not None and text_content != "":
                            file_contents["content"] = text_content
                            file_contents["content_type"] = "text"

                elif self._is_image(filename):
                    file_contents["image_data"] = file_content
                    file_contents["content_type"] = "image"

                else:
                    text_content = self._extract_text_content(filename, file_content)
                    if text_content is not None and text_content != "":
                        file_contents["content"] = text_content
                        file_contents["content_type"] = "text"
            except Exception as e:
                print(f"Error downloading file {filename}: {e}")
                return {}

        return file_contents

    def _is_image(self, filename: str) -> bool:
        return any(filename.lower().endswith(ext) for ext in self.image_extensions)

    def _is_pdf_image(self, file_content: bytes) -> bool:
        try:
            pdf = PdfReader(io.BytesIO(file_content))
            for page in pdf.pages:
                if "/XObject" in page["/Resources"]:
                    xObject = page["/Resources"]["/XObject"].get_object()
                    if isinstance(xObject, DictionaryObject):
                        for obj in xObject:
                            if xObject[obj].get_object()["/Subtype"] == "/Image":
                                return True
        except Exception as e:
            return False
        return False

    def retrieve_all_immigrant_files(
        self, immigrant_username: str
    ) -> dict[str, list[str]]:
        immigrant_files: dict[str, list[str]] = {}
        file_subfolders = ["ai_chat_files", "forms", "personal_files", "filled_forms"]
        for subfolder in file_subfolders:
            immigrant_files[subfolder] = []
            response = self.s3_client.list_objects_v2(
                Bucket=self.bucket_name,
                Prefix=f"immigrants/{immigrant_username.lower()}/{subfolder}/"
            )
            if 'Contents' in response:
                for obj in response["Contents"]:
                    file_key: str = obj['Key']
                    filename = file_key.split('/')[-1]
                    if filename and 'blob' not in filename and not self._is_image(filename):
                        immigrant_files[subfolder].append(f"{self.base_url}/{file_key}")
        return immigrant_files

    def retrieve_all_lawpersonnel_files(
        self, lawpersonnel_username: str, lawpersonnel_type: Literal["lawyers", "nonlawyers", "paralegals", "lawstudents"]
    ) -> dict[str, list[str]]:
        lawpersonnel_files: dict[str, list[str]] = {}
        file_subfolders = ["ai_chat_files", "personal_files", "profile_picture", "government_ids"]
        if lawpersonnel_type == "lawyers":
            file_subfolders.extend(
                ["intakes", "professional_licenses", "address_proofs"]
            )
        for subfolder in file_subfolders:
            lawpersonnel_files[subfolder] = []
            response = self.s3_client.list_objects_v2(
                Bucket=self.bucket_name,
                Prefix=f"{lawpersonnel_type}/{lawpersonnel_username.lower()}/{subfolder}/",
            )
            if "Contents" in response:
                for obj in response["Contents"]:
                    file_key: str = obj["Key"]
                    filename = file_key.split("/")[-1]
                    if (
                        filename
                        and "blob" not in filename
                        and not self._is_image(filename)
                    ):
                        lawpersonnel_files[subfolder].append(
                            f"{self.base_url}/{file_key}"
                        )
        return lawpersonnel_files

    def retrieve_all_case_files(self, immigrant_username: str, lawyer_map: dict[str, list[str]] = None) -> list[dict[str, Union[str, bytes]]]:
        all_immigrant_case_files = []
        directories = [f"immigrants/{immigrant_username.lower()}"]
        if lawyer_map:
            for lawyer_username, case_ids in lawyer_map.items():
                for case_id in case_ids:
                    directories.append(f"lawyers/{lawyer_username.lower()}/{case_id}/")

        paginator = self.s3_client.get_paginator('list_objects_v2')
        for directory in directories:
            for page in paginator.paginate(
                Bucket=self.bucket_name,
                Prefix=directory
            ):
                if 'Contents' in page:
                    for obj in page['Contents']:
                        file_key: str = obj["Key"]
                        filename, ext = os.path.splitext(os.path.basename(file_key))
                        if len(filename) > 0 and len(ext) > 0 and ext.lower() in self.allowed_case_extensions:
                            file_url = f"{self.base_url}/{file_key}"
                            file_contents = self.download_file(file_url)
                            file_contents["file_url"] = file_url
                            all_immigrant_case_files.append(file_contents)

        return all_immigrant_case_files

    def _create_dummy_file(self, filename: str, file_data: bytes) -> DummyFile:
        name, ext = os.path.splitext(filename)
        type, _ = mimetypes.guess_type(filename)
        return DummyFile(
            content=file_data,
            size=len(file_data),
            name=name,
            ext=ext,
            content_type=type
        )

    def add_lawyer_verification_items(self, lawyer_username: str, professional_license_filename: str, professional_license_contents: bytes, address_proof_filename: str, address_proof_contents: bytes) -> tuple[list[str], list[str]]:
        license_dir = f"lawyers/{lawyer_username}/professional_licenses"
        proof_dir = f"lawyers/{lawyer_username}/address_proofs"
        license_url = None
        proof_url = None
        proof_file = self._create_dummy_file(address_proof_filename, address_proof_contents)
        if not professional_license_filename.endswith("blob"):
            tmp_filename = Helpers.get_legal_filename(professional_license_filename)
            license_file = self._create_dummy_file(
                tmp_filename, professional_license_contents
            )
            s3_key, _ = self.crosscheck_existing(license_dir, tmp_filename)
            license_url = self.upload_file(license_file, s3_key)

        if not address_proof_filename.endswith("blob"):
            tmp_filename = Helpers.get_legal_filename(address_proof_filename)
            proof_file = self._create_dummy_file(tmp_filename, address_proof_contents)
            s3_key, _ = self.crosscheck_existing(proof_dir, tmp_filename)
            proof_url = self.upload_file(proof_file, s3_key)

        return license_url, proof_url

    def add_government_id(self, lawpersonnel_username: str, lawpersonnel_type: str, government_id_contents: bytes, government_id_filename: str) -> str:
        destination = f"{lawpersonnel_type}s/{lawpersonnel_username.lower()}/government_ids"
        if not government_id_filename.endswith("blob"):
            tmp_filename = Helpers.get_legal_filename(government_id_filename)
            government_id_file = self._create_dummy_file(tmp_filename, government_id_contents)
            s3_key, _ = self.crosscheck_existing(destination, tmp_filename)
            return self.upload_file(government_id_file, s3_key)

    def rename_directory(self, old_directory: str, new_directory: str) -> None:
        if old_directory.startswith(self.base_url):
            old_directory = old_directory.removeprefix(f"{self.base_url}/")
        elif old_directory.startswith("https"):
            return
        if new_directory.startswith(self.base_url):
            new_directory = new_directory.removeprefix(f"{self.base_url}/")
        elif new_directory.startswith("https"):
            return

        for obj in self.bucket.objects.filter(Prefix=old_directory):
            old_source = {"Bucket": self.bucket_name, "Key": obj.key}
            new_key = obj.key.replace(old_directory, new_directory, 1)
            new_obj = self.bucket.Object(new_key)
            new_obj.copy(old_source)

        for obj in self.bucket.objects.filter(Prefix=old_directory):
            self.s3_resource.Object(self.bucket_name, obj.key).delete()

    def create_directory(self, username: str, user_type: Literal["immigrants", "lawyers", "nonlawyers", "lawstudents", "paralegals"]):
        sub_folders = ["ai_chat_files", "personal_files", "profile_picture"]
        if user_type == "immigrants":
            sub_folders.extend(["forms", "filled_forms"])
        elif user_type == "lawyers":
            sub_folders.extend(["intakes", "professional_licenses", "address_proofs", "government_ids"])
        else:
            sub_folders.append("government_ids")

        for sub_folder in sub_folders:
            self.bucket.put_object(
                Key=f"{user_type.lower()}/{username.lower()}/{sub_folder}/",
                Body=b"",
                ACL="public-read"
            )

    def delete_user_directory(self, directory_url: str) -> None:
        if directory_url.startswith(self.base_url):
            directory_url = directory_url.removeprefix(f"{self.base_url}/")
        elif directory_url.startswith("https"):
            return

        objects_to_delete = [{"Key": obj.key} for obj in self.bucket.objects.filter(Prefix=directory_url)]
        if len(objects_to_delete) > 0:
            self.s3_client.delete_objects(
                Bucket=self.bucket_name,
                Delete={"Objects": objects_to_delete}
            )

    def delete_user_file(self, file_url: str) -> None:
        if file_url.startswith(self.base_url):
            file_url = file_url.removeprefix(f"{self.base_url}/")
        elif file_url.startswith("https"):
            return

        self.s3_client.delete_object(
            Bucket=self.bucket_name,
            Key=file_url
        )

    def move_user_file(self, file_url: str, new_directory: str) -> None:
        if file_url.startswith(self.base_url):
            file_url = file_url.removeprefix(f"{self.base_url}/")
        elif file_url.startswith("https"):
            return

        filename = file_url.split("/")[-1]
        self.s3_client.put_object(
            Bucket=self.bucket_name,
            Key=new_directory
        )
        new_file_url = f"{new_directory}/{filename}"
        self.s3_client.copy_object(
            Bucket=self.bucket_name,
            CopySource={"Bucket": self.bucket_name, "Key": file_url},
            Key=new_file_url
        )
        self.s3_client.delete_object(Bucket=self.bucket_name, Key=file_url)

    def upload_web_file(self, web_url: str, file_key: str) -> str:
        resp = requests.get(web_url, stream=True)
        self.s3_client.upload_fileobj(
            Fileobj=resp.raw,
            Bucket=self.bucket_name,
            Key=file_key,
            ExtraArgs={"ACL": "public-read"},
        )
        return f"{self.base_url}/{file_key}"
